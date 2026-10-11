
import re
import sys
import time
import unicodedata
from pathlib import Path
from urllib.parse import urljoin, urlparse, parse_qs, urlencode, urlunparse

import requests
from bs4 import BeautifulSoup
from PIL import Image, ImageTk
from io import BytesIO

import tkinter as tk
from tkinter import ttk, messagebox

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


BASE_URL = "https://www.db.yugioh-card.com/yugiohdb/"
PRODUCT_LIST_URL = urljoin(BASE_URL, "card_list.action?clm=3")

RARITIES = [
    "Common",
    "Short Print",
    "Rare",
    "Super Rare",
    "Ultra Rare",
    "Ultimate Rare",
    "Secret Rare",
    "Prismatic Secret Rare",
    "Collector's Rare",
    "Ghost Rare",
    "Gold Rare",
    "Gold Secret Rare",
    "Platinum Rare",
    "Platinum Secret Rare",
    "Starlight Rare",
    "Quarter Century Secret Rare",
    "Extra Secret Rare",
    "Parallel Rare",
]

RARITY_ABBREVIATIONS = {
    "C": "Common",
    "SP": "Short Print",
    "R": "Rare",
    "SR": "Super Rare",
    "UR": "Ultra Rare",
    "UL": "Ultimate Rare",
    "SE": "Secret Rare",
    "PS": "Prismatic Secret Rare",
    "CR": "Collector's Rare",
    "GR": "Ghost Rare",
    "GUR": "Gold Rare",
    "GS": "Gold Secret Rare",
    "P": "Parallel Rare",
    "STAR": "Starlight Rare",
    "QCSE": "Quarter Century Secret Rare",
    "EXSE": "Extra Secret Rare",
}

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/131.0 Safari/537.36"
    ),
    "Accept-Language": "en-US,en;q=0.9",
}

session = requests.Session()
session.headers.update(HEADERS)


def normalize(text):
    """Normalize names for case-insensitive expansion matching."""
    text = unicodedata.normalize("NFKD", text)
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def fetch_soup(url):
    response = session.get(url, timeout=30)
    response.raise_for_status()
    return BeautifulSoup(response.text, "html.parser")


def english_url(url):
    """Force the English locale on a database URL."""
    absolute = urljoin(BASE_URL, url)
    parsed = urlparse(absolute)
    query = parse_qs(parsed.query)
    query["request_locale"] = ["en"]
    return urlunparse(
        parsed._replace(query=urlencode(query, doseq=True))
    )


def find_expansion(expansion_name):
    """
    Find an expansion's official product-listing URL.

    The site exposes its products through card_list.action and
    individual product pages through card_search.action?ope=1&pid=...
    """
    print("Searching the official database for the expansion...")

    soup = fetch_soup(PRODUCT_LIST_URL)
    target = normalize(expansion_name)
    candidates = []

    product_entries = []
    for link in soup.select("a[href]"):
        href = link.get("href", "")
        label = " ".join(link.stripped_strings)
        if label and "card_search.action" in href:
            product_entries.append((label, href))

    for entry in soup.select('input.link_value[value*="card_search.action"]'):
        href = entry.get("value", "")
        label_node = entry.find_previous("p")
        label = " ".join(label_node.stripped_strings) if label_node else ""
        if label:
            product_entries.append((label, href))

    for label, href in product_entries:
        query = parse_qs(urlparse(urljoin(BASE_URL, href)).query)
        if query.get("ope", [""])[0] != "1":
            continue
        if not query.get("pid"):
            continue

        candidates.append((label, english_url(href)))

    # Exact match first.
    for label, url in candidates:
        if normalize(label) == target:
            return label, url

    # Permit an expansion name with minor punctuation differences.
    matches = [
        (label, url)
        for label, url in candidates
        if target in normalize(label)
        or normalize(label) in target
    ]

    if len(matches) == 1:
        return matches[0]

    if len(matches) > 1:
        exact_word_matches = [
            item for item in matches
            if normalize(item[0]) == target
        ]
        if len(exact_word_matches) == 1:
            return exact_word_matches[0]

        choices = "\n".join(f"  - {label}" for label, _ in matches[:15])
        raise ValueError(
            "Multiple expansions matched your input:\n"
            f"{choices}\n\nPlease enter a more specific name."
        )

    raise ValueError(
        f'Expansion "{expansion_name}" was not found on the '
        "currently accessible product-listing page. Check the spelling "
        "and try again. Older expansions may require browsing the "
        "database's product listings."
    )


def find_card_links(product_url):
    """Collect unique card detail URLs from an expansion page."""
    soup = fetch_soup(product_url)
    cards = []
    seen = set()

    card_links = [
        link.get("href", "")
        for link in soup.select('a[href*="card_search.action"]')
    ]
    card_links.extend(
        entry.get("value", "")
        for entry in soup.select('input.link_value[value*="card_search.action"]')
    )

    for href in card_links:
        absolute = urljoin(BASE_URL, href)
        query = parse_qs(urlparse(absolute).query)

        if query.get("ope", [""])[0] != "2":
            continue
        if not query.get("cid"):
            continue

        card_id = query["cid"][0]
        if card_id in seen:
            continue

        seen.add(card_id)
        cards.append(english_url(absolute))

    if not cards:
        raise ValueError(
            "No card detail links were found in the expansion page. "
            "The website's HTML structure may have changed."
        )

    return cards


def text_without_noise(element):
    """Extract readable card text without scripts and styles."""
    if element is None:
        return ""

    for node in element.select("script, style, noscript"):
        node.decompose()

    return " ".join(element.stripped_strings)


def find_card_name(soup):
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    detail_title = re.match(r"^(.*?)\s*\|\s*Card Details\b", title, re.I)
    if detail_title and detail_title.group(1).strip():
        return detail_title.group(1).strip()

    for node in soup.select("h1"):
        value = " ".join(node.stripped_strings)
        if value and "trading card game card database" not in normalize(value):
            return value

    return re.sub(r"\s*\|.*$", "", title).strip() or "Unknown card"


def find_card_text(soup):
    """Extract the card effect/description from a detail page."""
    for selector in [
        "#card_text",
        ".card_text",
        ".item_box_text",
        ".box_card_text",
    ]:
        node = soup.select_one(selector)
        if node:
            value = text_without_noise(node)
            if value:
                return value

    # Fallback: locate the Card Text heading and read the following block.
    for node in soup.find_all(
        lambda tag: tag.name in ("h2", "h3", "dt", "strong")
        and tag.get_text(" ", strip=True).lower() == "card text"
    ):
        following = node.find_next(["p", "dd", "div"])
        if following:
            value = text_without_noise(following)
            if value:
                return value

    return "Card text could not be extracted automatically."


def find_card_image(soup):
    """Find the illustration image on the card detail page."""
    # Prefer the page's primary image when available.
    for selector in [
        'img[id*="card_image"]',
        'img[class*="card_image"]',
        ".card_image img",
        ".image_set img",
        'meta[property="og:image"]',
    ]:
        node = soup.select_one(selector)
        if not node:
            continue

        src = node.get("content") or node.get("src") or node.get("data-src")
        if src and not src.startswith("data:"):
            return urljoin(BASE_URL, src)

    # Fallback: look for an image URL that resembles a card image.
    for node in soup.select("img[src], img[data-src]"):
        src = node.get("data-src") or node.get("src", "")
        if not src or src.startswith("data:"):
            continue
        if any(word in src.lower() for word in ("card", "image", "monster")):
            return urljoin(BASE_URL, src)

    return None


def extract_set_number(soup, expansion_name):
    """
    Find an English TCG set number such as BETB-EN081.

    If multiple set numbers appear, prefer the one nearest the requested
    expansion's name; otherwise use the first English set number.
    """
    text = soup.get_text(" ", strip=True)
    matches = list(
        re.finditer(r"\b[A-Z0-9]{2,8}-EN\d{2,4}\b", text, re.I)
    )

    if not matches:
        return ""

    target = normalize(expansion_name)
    expansion_position = text.lower().find(expansion_name.lower())

    if expansion_position >= 0:
        nearest = min(
            matches,
            key=lambda match: abs(match.start() - expansion_position),
        )
        return nearest.group(0).upper()

    return matches[0].group(0).upper()


def extract_rarities(soup, expansion_name):
    """
    Extract rarities from this card's printings in the requested expansion.
    """
    found = []

    for row in soup.select("#update_list .t_row"):
        pack_name = row.select_one(".pack_name")
        if not pack_name or normalize(" ".join(pack_name.stripped_strings)) != normalize(expansion_name):
            continue

        for rarity_node in row.select(".icon.rarity span"):
            rarity = " ".join(rarity_node.stripped_strings)
            if rarity in RARITIES:
                found.append(rarity)

    # Preserve the order defined in RARITIES.
    found_set = set(found)
    return [rarity for rarity in RARITIES if rarity in found_set]


def get_card_data(url, expansion_name):
    soup = fetch_soup(url)

    name = find_card_name(soup)
    description = find_card_text(soup)
    image_url = find_card_image(soup)
    card_number = extract_set_number(soup, expansion_name)
    rarities = extract_rarities(soup, expansion_name)

    return {
        "name": name,
        "description": description,
        "image_url": image_url,
        "card_number": card_number,
        "rarities": rarities,
        "url": url,
    }


def safe_filename(name):
    """Convert an expansion name into a Windows-safe .xlsx filename."""
    filename = normalize(name).replace(" ", "_")
    return filename or "wishlist"


def wishlist_path(expansion_name):
    return Path(__file__).resolve().parent / (
        safe_filename(expansion_name) + ".xlsx"
    )


def parse_saved_quantity(value):
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return max(value, 0)
    if isinstance(value, float):
        return int(value) if value >= 0 and value.is_integer() else 0

    text = str(value or "").strip()
    checkbox_symbols = ("☐", "☑", "□", "☒", "✓", "✔")
    checkbox_count = sum(text.count(symbol) for symbol in checkbox_symbols)
    if checkbox_count:
        return checkbox_count
    try:
        return max(0, int(text))
    except ValueError:
        return 0


def load_existing_quantities(expansion_name, cards):
    """Load quantities from a previously exported wishlist, if present."""
    path = wishlist_path(expansion_name)
    if not path.exists():
        return None

    quantities = {
        card["url"]: {rarity: 0 for rarity in card["rarities"]}
        for card in cards
    }
    by_number = {}
    by_name = {}
    for card in cards:
        if card["card_number"]:
            by_number.setdefault(card["card_number"].casefold(), []).append(card)
        by_name.setdefault(normalize(card["name"]), []).append(card)

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Wishlist"] if "Wishlist" in workbook.sheetnames else workbook.active
        for quantity_value, card_name, rarity, card_number in sheet.iter_rows(
            min_row=2, min_col=1, max_col=4, values_only=True
        ):
            if not rarity:
                continue

            matches = by_number.get(str(card_number).strip().casefold(), []) if card_number else []
            if not matches and card_name:
                matches = by_name.get(normalize(str(card_name)), [])

            for card in matches:
                if rarity in quantities[card["url"]]:
                    quantities[card["url"]][rarity] += parse_saved_quantity(
                        quantity_value
                    )
    finally:
        workbook.close()

    return quantities


def write_excel(expansion_name, selected_cards):
    """Create the wishlist workbook using only positive quantities."""
    filename = wishlist_path(expansion_name)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Wishlist"

    headers = ["Cantidad", "Nombre", "Rareza", "# de carta"]
    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for card in selected_cards:
        for rarity, quantity in card["quantities"].items():
            if quantity <= 0:
                continue

            # Represent each requested copy as an empty checkbox.
            checkboxes = "".join(["☐"] * quantity)

            sheet.append([
                checkboxes,
                card["name"],
                rarity,
                card["card_number"],
            ])

    widths = {
        "A": 24,
        "B": 42,
        "C": 30,
        "D": 18,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    workbook.save(filename)
    return filename


class WishlistApp:
    def __init__(self, root, expansion_name, cards, initial_quantities=None):
        self.root = root
        self.expansion_name = expansion_name
        self.cards = cards
        self.index = 0
        self.quantities = {}
        self.photo = None
        self.goto_var = tk.StringVar(value="1")

        for card in self.cards:
            self.quantities[card["url"]] = {
                rarity: 0 for rarity in card["rarities"]
            }
            if initial_quantities and card["url"] in initial_quantities:
                self.quantities[card["url"]].update(
                    initial_quantities[card["url"]]
                )

        self.root.title(f"Yu-Gi-Oh! Wishlist — {expansion_name}")
        self.root.geometry("1050x800")
        self.root.minsize(780, 620)

        self.build_ui()
        self.show_card()

    def build_ui(self):
        self.root.columnconfigure(0, weight=1)
        self.root.rowconfigure(1, weight=1)

        header = ttk.Frame(self.root, padding=10)
        header.grid(row=0, column=0, sticky="ew")
        header.columnconfigure(0, weight=1)

        self.title_label = ttk.Label(
            header,
            text=self.expansion_name,
            font=("Segoe UI", 17, "bold"),
        )
        self.title_label.grid(row=0, column=0, sticky="w")

        self.counter_label = ttk.Label(header, text="")
        self.counter_label.grid(row=0, column=1, sticky="e")

        content = ttk.Frame(self.root, padding=(12, 0, 12, 8))
        content.grid(row=1, column=0, sticky="nsew")
        content.columnconfigure(0, weight=1)
        content.columnconfigure(1, weight=2)
        content.rowconfigure(0, weight=1)

        # Illustration panel.
        image_frame = ttk.LabelFrame(
            content, text="Illustration", padding=10
        )
        image_frame.grid(
            row=0, column=0, sticky="nsew", padx=(0, 8)
        )
        image_frame.columnconfigure(0, weight=1)
        image_frame.rowconfigure(0, weight=1)

        self.image_label = ttk.Label(
            image_frame,
            text="Loading illustration...",
            anchor="center",
            justify="center",
        )
        self.image_label.grid(row=0, column=0, sticky="nsew")

        # Card information panel.
        info_frame = ttk.Frame(content)
        info_frame.grid(row=0, column=1, sticky="nsew")
        info_frame.columnconfigure(0, weight=1)
        info_frame.rowconfigure(2, weight=1)

        self.name_label = ttk.Label(
            info_frame,
            text="",
            font=("Segoe UI", 15, "bold"),
            wraplength=590,
        )
        self.name_label.grid(row=0, column=0, sticky="w", pady=(0, 8))

        self.number_label = ttk.Label(info_frame, text="")
        self.number_label.grid(row=1, column=0, sticky="w", pady=(0, 8))

        text_frame = ttk.LabelFrame(
            info_frame, text="Card Text", padding=6
        )
        text_frame.grid(row=2, column=0, sticky="nsew", pady=(0, 10))
        text_frame.columnconfigure(0, weight=1)
        text_frame.rowconfigure(0, weight=1)

        self.description = tk.Text(
            text_frame,
            wrap="word",
            height=12,
            width=45,
            font=("Segoe UI", 10),
            state="disabled",
        )
        text_scroll = ttk.Scrollbar(
            text_frame, command=self.description.yview
        )
        self.description.configure(yscrollcommand=text_scroll.set)
        self.description.grid(row=0, column=0, sticky="nsew")
        text_scroll.grid(row=0, column=1, sticky="ns")

        rarity_frame = ttk.LabelFrame(
            info_frame, text="Desired Copies by Rarity", padding=8
        )
        rarity_frame.grid(row=3, column=0, sticky="ew")
        rarity_frame.columnconfigure(0, weight=1)

        self.rarity_rows = ttk.Frame(rarity_frame)
        self.rarity_rows.grid(row=0, column=0, sticky="ew")

        self.rarity_vars = {}

        # Navigation and completion buttons.
        footer = ttk.Frame(self.root, padding=12)
        footer.grid(row=2, column=0, sticky="ew")
        footer.columnconfigure(4, weight=1)

        self.previous_button = ttk.Button(
            footer, text="← Previous", command=self.previous_card
        )
        self.previous_button.grid(row=0, column=0, sticky="w")

        validate_card_number = self.root.register(self.validate_card_number)
        self.goto_entry = ttk.Entry(
            footer,
            textvariable=self.goto_var,
            width=6,
            validate="key",
            validatecommand=(validate_card_number, "%P"),
        )
        self.goto_entry.grid(row=0, column=3, sticky="w")

        self.goto_button = ttk.Button(
            footer, text="Go to", command=self.go_to_card
        )
        self.goto_button.grid(row=0, column=4, sticky="w", padx=(4, 0))

        self.next_button = ttk.Button(
            footer, text="Next →", command=self.next_card
        )
        self.next_button.grid(row=0, column=1, sticky="w", padx=(8, 8))

        self.finish_button = ttk.Button(
            footer, text="Finish and Export Excel",
            command=self.finish,
        )
        self.finish_button.grid(row=0, column=5, sticky="e")

    def save_current_quantities(self):
        if not self.cards:
            return

        card = self.cards[self.index]
        saved = self.quantities[card["url"]]

        for rarity, variable in self.rarity_vars.items():
            saved[rarity] = variable.get()

    def show_card(self):
        card = self.cards[self.index]

        self.counter_label.configure(
            text=f"Card {self.index + 1} of {len(self.cards)}"
        )
        self.goto_var.set(str(self.index + 1))
        self.name_label.configure(text=card["name"])
        self.number_label.configure(
            text=f"Card number: {card['card_number'] or 'Not found'}"
        )

        self.description.configure(state="normal")
        self.description.delete("1.0", "end")
        self.description.insert("1.0", card["description"])
        self.description.configure(state="disabled")

        # Rebuild the rarity controls for the current card.
        for child in self.rarity_rows.winfo_children():
            child.destroy()

        self.rarity_vars = {}
        saved = self.quantities[card["url"]]

        if not card["rarities"]:
            ttk.Label(
                self.rarity_rows,
                text="No rarities were detected automatically.",
                wraplength=450,
            ).grid(row=0, column=0, sticky="w")

        for row, rarity in enumerate(card["rarities"]):
            ttk.Label(
                self.rarity_rows, text=rarity
            ).grid(row=row, column=0, sticky="w", pady=2)

            variable = tk.IntVar(value=saved.get(rarity, 0))
            spinbox = ttk.Spinbox(
                self.rarity_rows,
                from_=0,
                to=999,
                increment=1,
                width=6,
                textvariable=variable,
                command=lambda r=rarity, v=variable: self.update_quantity(
                    r, v
                ),
            )
            spinbox.grid(row=row, column=1, sticky="e", padx=(12, 0))
            spinbox.bind(
                "<FocusOut>",
                lambda event, r=rarity, v=variable: self.update_quantity(
                    r, v
                ),
            )
            spinbox.bind(
                "<Return>",
                lambda event, r=rarity, v=variable: self.update_quantity(
                    r, v
                ),
            )
            self.rarity_vars[rarity] = variable

        self.load_illustration(card["image_url"])

    @staticmethod
    def validate_card_number(value):
        return value == "" or re.fullmatch(r"[0-9]*", value) is not None

    def go_to_card(self):
        value = self.goto_var.get()
        if not value:
            messagebox.showwarning(
                "Invalid card number",
                f"Enter a card number from 1 to {len(self.cards)}.",
            )
            return

        target = int(value)
        if not 1 <= target <= len(self.cards):
            messagebox.showwarning(
                "Invalid card number",
                f"Enter a card number from 1 to {len(self.cards)}.",
            )
            return

        self.save_current_quantities()
        self.index = target - 1
        self.show_card()

    def update_quantity(self, rarity, variable):
        try:
            value = max(0, int(variable.get()))
        except (ValueError, tk.TclError):
            value = 0

        card = self.cards[self.index]
        self.quantities[card["url"]][rarity] = value

    def load_illustration(self, image_url):
        self.photo = None
        self.image_label.configure(
            image="", text="Loading illustration..."
        )
        self.root.update_idletasks()

        if not image_url:
            self.image_label.configure(
                text="Illustration unavailable"
            )
            return

        try:
            response = session.get(image_url, timeout=30)
            response.raise_for_status()

            image = Image.open(BytesIO(response.content)).convert("RGB")
            image.thumbnail((390, 570), Image.Resampling.LANCZOS)

            self.photo = ImageTk.PhotoImage(image)
            self.image_label.configure(image=self.photo, text="")

        except Exception as exc:
            self.image_label.configure(
                image="",
                text=f"Could not load illustration:\n{exc}",
            )

    def previous_card(self):
        self.save_current_quantities()
        self.index = (self.index - 1) % len(self.cards)
        self.show_card()

    def next_card(self):
        self.save_current_quantities()
        self.index = (self.index + 1) % len(self.cards)
        self.show_card()

    def finish(self):
        self.save_current_quantities()

        selected_cards = []
        for card in self.cards:
            quantities = self.quantities[card["url"]]
            if any(value > 0 for value in quantities.values()):
                selected_cards.append({
                    **card,
                    "quantities": quantities.copy(),
                })

        if not selected_cards:
            messagebox.showwarning(
                "Empty wishlist",
                "You have not selected any cards. "
                "Set at least one rarity quantity above zero.",
            )
            return

        try:
            output = write_excel(self.expansion_name, selected_cards)
            messagebox.showinfo(
                "Wishlist created",
                f"Your Excel wishlist was saved to:\n{output}",
            )
        except Exception as exc:
            messagebox.showerror(
                "Export failed",
                f"Could not create the Excel file:\n{exc}",
            )


def main():
    if len(sys.argv) > 1:
        expansion_name = " ".join(sys.argv[1:]).strip()
    else:
        # Simple prompt before the main GUI opens.
        prompt_root = tk.Tk()
        prompt_root.withdraw()

        from tkinter import simpledialog

        expansion_name = simpledialog.askstring(
            "Yu-Gi-Oh! Wishlist",
            "Enter the expansion name (e.g. Beyond the Brave):",
            parent=prompt_root,
        )
        prompt_root.destroy()

        if not expansion_name or not expansion_name.strip():
            return

        expansion_name = expansion_name.strip()

    try:
        _, product_url = find_expansion(expansion_name)
        print(f"Expansion page: {product_url}")

        card_urls = find_card_links(product_url)
        print(f"Found {len(card_urls)} card detail pages.")

        cards = []
        for index, url in enumerate(card_urls, start=1):
            print(f"Loading card {index}/{len(card_urls)}...")
            try:
                card = get_card_data(url, expansion_name)
                cards.append(card)
            except requests.RequestException as exc:
                print(f"Skipping unavailable card page: {url}: {exc}")

            # Be considerate of the official website.
            time.sleep(0.15)

        if not cards:
            raise ValueError(
                "No card details could be loaded successfully."
            )

    except Exception as exc:
        error_root = tk.Tk()
        error_root.withdraw()
        messagebox.showerror("Could not load expansion", str(exc))
        error_root.destroy()
        return

    existing_quantities = None
    workbook_error = None
    try:
        existing_quantities = load_existing_quantities(expansion_name, cards)
    except Exception as exc:
        workbook_error = exc

    root = tk.Tk()
    WishlistApp(root, expansion_name, cards, existing_quantities)
    if workbook_error:
        messagebox.showwarning(
            "Could not read existing wishlist",
            f"The existing Excel file could not be loaded:\n{workbook_error}",
            parent=root,
        )
    root.mainloop()


if __name__ == "__main__":
    main()
